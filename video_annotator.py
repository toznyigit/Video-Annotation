# ==============================================================================
# Video Annotation Tool
# Copyright (C) 2026 Tamer Yigit and contributors
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.
# ==============================================================================

import sys
import os
import csv
from datetime import timedelta

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QFileDialog, QLineEdit, QComboBox,
    QTableWidget, QTableWidgetItem, QSplitter, QGroupBox,
    QSpinBox, QDoubleSpinBox, QTextEdit, QMessageBox, QHeaderView,
    QDialog, QTextBrowser,
)
from PyQt6.QtMultimedia import QMediaPlayer, QAudioOutput
from PyQt6.QtMultimediaWidgets import QVideoWidget
from PyQt6.QtCore import Qt, QUrl, QTimer, QEvent
from PyQt6.QtGui import QFont, QColor, QDesktopServices

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

__version__ = "1.0.0"

# ── Colour palette ─────────────────────────────────────────────────────────────
DARK_BG    = "#1e1e2e"
PANEL_BG   = "#2a2a3e"
ACCENT     = "#7c6af7"
ACCENT2    = "#56cfb2"
MOTHER_CLR = "#f4a261"
CHILD_CLR  = "#56cfb2"
TEXT_MAIN  = "#e0e0f0"
TEXT_MUTED = "#8888aa"
BTN_NEUTRAL= "#4a4a6a"

STYLESHEET = f"""
QMainWindow, QWidget {{
    background-color: {DARK_BG};
    color: {TEXT_MAIN};
    font-family: 'Segoe UI', Arial, sans-serif;
    font-size: 13px;
}}
QGroupBox {{
    border: 1px solid {ACCENT};
    border-radius: 8px;
    margin-top: 10px;
    padding-top: 10px;
    font-weight: bold;
    color: {ACCENT};
}}
QGroupBox::title {{
    subcontrol-origin: margin;
    left: 12px;
    padding: 0 6px;
}}
QPushButton {{
    background-color: {ACCENT};
    color: white;
    border: none;
    border-radius: 6px;
    padding: 7px 16px;
    font-weight: bold;
}}
QPushButton:hover  {{ background-color: #9b8cf9; }}
QPushButton:pressed {{ background-color: #5a4bc0; }}
QPushButton:disabled {{ background-color: {BTN_NEUTRAL}; color: {TEXT_MUTED}; }}
QPushButton#success {{ background-color: #2ecc71; }}
QPushButton#success:hover {{ background-color: #27ae60; }}
QPushButton#danger  {{ background-color: #e74c3c; }}
QPushButton#danger:hover  {{ background-color: #c0392b; }}
QPushButton#neutral {{ background-color: {BTN_NEUTRAL}; }}
QPushButton#neutral:hover {{ background-color: #5a5a8a; }}
QLineEdit, QSpinBox, QDoubleSpinBox, QComboBox, QTextEdit, QTextBrowser {{
    background-color: {PANEL_BG};
    color: {TEXT_MAIN};
    border: 1px solid {ACCENT};
    border-radius: 5px;
    padding: 5px 8px;
}}
QLineEdit:focus, QSpinBox:focus, QDoubleSpinBox:focus,
QComboBox:focus, QTextEdit:focus, QTextBrowser:focus {{
    border-color: {ACCENT2};
}}
QComboBox::drop-down {{ border: none; }}
QComboBox QAbstractItemView {{
    background-color: {PANEL_BG};
    color: {TEXT_MAIN};
    selection-background-color: {ACCENT};
}}
QTableWidget {{
    background-color: {PANEL_BG};
    color: {TEXT_MAIN};
    gridline-color: #3a3a5a;
    border: 1px solid {ACCENT};
    border-radius: 6px;
    alternate-background-color: #252540;
}}
QTableWidget::item:selected {{
    background-color: {ACCENT};
    color: white;
}}
QHeaderView::section {{
    background-color: #35355a;
    color: {TEXT_MAIN};
    border: none;
    padding: 6px;
    font-weight: bold;
}}
QScrollBar:vertical {{
    background: {PANEL_BG};
    width: 10px;
    border-radius: 5px;
}}
QScrollBar::handle:vertical {{
    background: {ACCENT};
    border-radius: 5px;
    min-height: 20px;
}}
QSplitter::handle {{
    background-color: {ACCENT};
    width: 2px;
}}
"""


def fmt_time(ms: int) -> str:
    s = ms // 1000
    return str(timedelta(seconds=s))


class VideoAnnotator(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"Video Annotation Tool v{__version__} [GPL-3.0]")
        self.resize(1400, 880)

        self.video_path  = ""
        self.questions   = []
        self.records     = []
        self.current_q   = 0
        self.answers_buf = {}
        self.step_ms     = 1000
        self.range_start = 0
        self.range_end   = 0
        self.label       = "Mother"
        self.processing  = False
        self.waiting_ans = False
        self.current_pos = 0
        self.total_steps = 0

        self._build_ui()
        self.setStyleSheet(STYLESHEET)

    # ── UI ─────────────────────────────────────────────────────────────────────
    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setSpacing(8)
        root.setContentsMargins(10, 8, 10, 10)

        # Top Bar
        top_bar = QHBoxLayout()
        title_lbl = QLabel("🎬 Video Annotation Tool")
        title_lbl.setFont(QFont("Segoe UI", 13, QFont.Weight.Bold))
        title_lbl.setStyleSheet(f"color: {ACCENT2};")
        ver_lbl = QLabel(f"v{__version__} • GPL-3.0")
        ver_lbl.setStyleSheet(f"color: {TEXT_MUTED}; font-size: 11px;")

        top_bar.addWidget(title_lbl)
        top_bar.addWidget(ver_lbl)
        top_bar.addStretch()

        btn_manual = QPushButton("📖 User Manual")
        btn_manual.setObjectName("neutral")
        btn_manual.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_manual.clicked.connect(self._show_manual)
        top_bar.addWidget(btn_manual)

        btn_about = QPushButton("⚖️ About / License")
        btn_about.setObjectName("neutral")
        btn_about.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_about.clicked.connect(self._show_about)
        top_bar.addWidget(btn_about)

        root.addLayout(top_bar)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        root.addWidget(splitter, stretch=1)
        splitter.addWidget(self._build_left())
        splitter.addWidget(self._build_right())
        splitter.setSizes([680, 720])

    def _build_left(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setSpacing(8)

        self.video_widget = QVideoWidget()
        self.video_widget.setMinimumHeight(340)
        self.video_widget.setStyleSheet("background: #000;")
        lay.addWidget(self.video_widget)

        self.player = QMediaPlayer()
        self.audio  = QAudioOutput()
        self.player.setAudioOutput(self.audio)
        self.player.setVideoOutput(self.video_widget)
        self.player.positionChanged.connect(self._on_position_changed)
        self.player.durationChanged.connect(self._on_duration_changed)

        self.time_lbl = QLabel("00:00:00 / 00:00:00")
        self.time_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.time_lbl.setFont(QFont("Courier New", 12, QFont.Weight.Bold))
        self.time_lbl.setStyleSheet(f"color: {ACCENT2};")
        lay.addWidget(self.time_lbl)

        # File loading group
        fg = QGroupBox("Load Files")
        fl = QVBoxLayout(fg)
        vrow = QHBoxLayout()
        self.video_path_lbl = QLineEdit(); self.video_path_lbl.setPlaceholderText("No video selected…"); self.video_path_lbl.setReadOnly(True)
        btn_vid = QPushButton("Browse Video"); btn_vid.clicked.connect(self._browse_video)
        vrow.addWidget(self.video_path_lbl); vrow.addWidget(btn_vid)
        fl.addLayout(vrow)
        qrow = QHBoxLayout()
        self.txt_path_lbl = QLineEdit(); self.txt_path_lbl.setPlaceholderText("No questions file…"); self.txt_path_lbl.setReadOnly(True)
        btn_txt = QPushButton("Browse Questions"); btn_txt.clicked.connect(self._browse_txt)
        qrow.addWidget(self.txt_path_lbl); qrow.addWidget(btn_txt)
        fl.addLayout(qrow)
        lay.addWidget(fg)

        # Settings group
        cg = QGroupBox("Session Settings")
        cl = QVBoxLayout(cg)
        rng = QHBoxLayout()
        rng.addWidget(QLabel("Start (s):")); self.start_spin = QSpinBox(); self.start_spin.setRange(0,999999); rng.addWidget(self.start_spin)
        rng.addWidget(QLabel("End (s):")); self.end_spin = QSpinBox(); self.end_spin.setRange(0,999999); self.end_spin.setValue(60); rng.addWidget(self.end_spin)
        cl.addLayout(rng)
        sl = QHBoxLayout()
        sl.addWidget(QLabel("Step (s):")); self.step_spin = QDoubleSpinBox(); self.step_spin.setRange(0.1,300); self.step_spin.setValue(1.0); self.step_spin.setSingleStep(0.5); sl.addWidget(self.step_spin)
        sl.addWidget(QLabel("Label:"))
        self.label_combo = QComboBox(); self.label_combo.addItems(["Mother","Child"]); self.label_combo.currentTextChanged.connect(self._on_label_change); sl.addWidget(self.label_combo)
        self.badge = QLabel("● Mother"); self.badge.setStyleSheet(f"color:{MOTHER_CLR}; font-weight:bold;"); sl.addWidget(self.badge)
        cl.addLayout(sl)
        lay.addWidget(cg)

        # Action buttons
        br = QHBoxLayout()
        self.btn_start = QPushButton("▶  Start Processing"); self.btn_start.setObjectName("success"); self.btn_start.setMinimumHeight(38); self.btn_start.clicked.connect(self._start_processing)
        self.btn_stop  = QPushButton("■  Stop"); self.btn_stop.setObjectName("danger"); self.btn_stop.setEnabled(False); self.btn_stop.clicked.connect(self._stop_processing)
        br.addWidget(self.btn_start); br.addWidget(self.btn_stop)
        lay.addLayout(br)

        self.status_lbl = QLabel("Ready.")
        self.status_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.status_lbl.setStyleSheet(f"color:{ACCENT2}; font-style:italic;")
        lay.addWidget(self.status_lbl)
        lay.addStretch()
        return w

    def _build_right(self):
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setSpacing(8)

        qa_grp = QGroupBox("Current Question")
        qg = QVBoxLayout(qa_grp)

        self.q_counter = QLabel("Question — / —")
        self.q_counter.setFont(QFont("Segoe UI", 10, QFont.Weight.Bold))
        qg.addWidget(self.q_counter)

        self.q_text = QLabel("(no active session)")
        self.q_text.setWordWrap(True)
        self.q_text.setFont(QFont("Segoe UI", 12))
        self.q_text.setStyleSheet(f"color:{ACCENT2}; padding:6px; background:{PANEL_BG}; border-radius:5px;")
        self.q_text.setMinimumHeight(50)
        qg.addWidget(self.q_text)

        self.ans_input = QTextEdit()
        self.ans_input.setPlaceholderText("Type your answer here… (Ctrl+Enter or Cmd+Enter to submit)")
        self.ans_input.setMaximumHeight(90)
        self.ans_input.setEnabled(False)
        self.ans_input.installEventFilter(self)
        qg.addWidget(self.ans_input)

        ab = QHBoxLayout()
        self.btn_next_q = QPushButton("Next Question  ›"); self.btn_next_q.clicked.connect(self._next_question); self.btn_next_q.setEnabled(False)
        self.btn_skip_q = QPushButton("Skip"); self.btn_skip_q.setObjectName("neutral"); self.btn_skip_q.clicked.connect(self._skip_question); self.btn_skip_q.setEnabled(False)
        ab.addWidget(self.btn_next_q); ab.addWidget(self.btn_skip_q)
        qg.addLayout(ab)
        lay.addWidget(qa_grp)

        tbl_grp = QGroupBox("Annotation Records")
        tg = QVBoxLayout(tbl_grp)
        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["Video","Label","Timestamp","Step #","Question","Answer"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        tg.addWidget(self.table)

        tb = QHBoxLayout()
        self.btn_export = QPushButton("⬇  Export to Excel"); self.btn_export.setObjectName("success"); self.btn_export.clicked.connect(self._export_excel); self.btn_export.setEnabled(False)
        btn_clear = QPushButton("🗑  Clear Table"); btn_clear.setObjectName("danger"); btn_clear.clicked.connect(self._clear_table)
        tb.addWidget(self.btn_export); tb.addWidget(btn_clear)
        tg.addLayout(tb)
        lay.addWidget(tbl_grp, stretch=1)
        return w

    # ── Event Filter (Keyboard Shortcuts) ──────────────────────────────────────
    def eventFilter(self, obj, event):
        if obj == self.ans_input and event.type() == QEvent.Type.KeyPress:
            if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter):
                # Ctrl+Enter or Cmd+Enter submits the question
                if event.modifiers() & (Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.MetaModifier):
                    if self.btn_next_q.isEnabled():
                        self._next_question()
                        return True
        return super().eventFilter(obj, event)

    # ── In-App Help & Dialogs ──────────────────────────────────────────────────
    def _show_manual(self):
        # Locate USER_MANUAL.md next to script or in PyInstaller bundle
        base_dir = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
        manual_path = os.path.join(base_dir, "USER_MANUAL.md")
        if not os.path.exists(manual_path):
            manual_path = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "USER_MANUAL.md")

        content = ""
        if os.path.exists(manual_path):
            try:
                with open(manual_path, "r", encoding="utf-8") as f:
                    content = f.read()
            except Exception as e:
                content = f"Error reading manual: {e}"
        else:
            content = (
                "# Video Annotation Tool — Quick Guide\n\n"
                "### Workflow:\n"
                "1. **Load Video & Questions**: Browse to select a video and a `.txt` file with one question per line.\n"
                "2. **Configure Range & Step**: Set start time, end time, and step interval (e.g. 1.0s).\n"
                "3. **Choose Label**: Select participant role (Mother or Child).\n"
                "4. **Click Start Processing**: At each step, type your answer and click 'Next Question' (or Ctrl+Enter).\n"
                "5. **Export to Excel**: Once finished, click 'Export to Excel' to save a styled `.xlsx` file.\n"
            )

        dlg = QDialog(self)
        dlg.setWindowTitle("User Manual — Video Annotation Tool")
        dlg.resize(900, 700)
        d_lay = QVBoxLayout(dlg)
        d_lay.setSpacing(10)
        d_lay.setContentsMargins(12, 12, 12, 12)

        browser = QTextBrowser()
        browser.setOpenExternalLinks(True)
        browser.setMarkdown(content)
        browser.setStyleSheet(f"background: {PANEL_BG}; color: {TEXT_MAIN}; border: 1px solid {ACCENT}; border-radius: 6px; padding: 10px;")
        d_lay.addWidget(browser)

        btn_row = QHBoxLayout()
        btn_row.addStretch()
        btn_close = QPushButton("Close")
        btn_close.clicked.connect(dlg.accept)
        btn_row.addWidget(btn_close)
        d_lay.addLayout(btn_row)

        dlg.exec()

    def _show_about(self):
        about_text = (
            f"<h3>Video Annotation Tool v{__version__}</h3>"
            "<p>A desktop video observation and interval annotation system.</p>"
            "<p><b>License:</b> GNU General Public License v3.0 (GPL-3.0)<br>"
            "This is free software; you are free to change and redistribute it under GPLv3.<br>"
            "There is NO WARRANTY, to the extent permitted by law.</p>"
            "<p>Full license terms are available in the bundled <code>LICENSE</code> file or at:<br>"
            "<a href='https://www.gnu.org/licenses/gpl-3.0.html' style='color:#56cfb2;'>https://www.gnu.org/licenses/gpl-3.0.html</a></p>"
        )
        QMessageBox.about(self, "About Video Annotation Tool", about_text)

    # ── File browsing ──────────────────────────────────────────────────────────
    def _browse_video(self):
        path, _ = QFileDialog.getOpenFileName(self,"Open Video","","Video Files (*.mp4 *.avi *.mkv *.mov *.wmv *.flv *.webm)")
        if path:
            self.video_path = path
            self.video_path_lbl.setText(path)
            self.player.setSource(QUrl.fromLocalFile(path))
            self.player.pause()
            self._set_status(f"Loaded: {os.path.basename(path)}")

    def _browse_txt(self):
        path, _ = QFileDialog.getOpenFileName(self,"Open Questions File","","Text Files (*.txt)")
        if path:
            with open(path,"r",encoding="utf-8") as f:
                self.questions = [l.strip() for l in f if l.strip()]
            self.txt_path_lbl.setText(path)
            self._set_status(f"Loaded {len(self.questions)} questions.")

    def _on_label_change(self, text):
        self.label = text
        clr = MOTHER_CLR if text == "Mother" else CHILD_CLR
        self.badge.setText(f"● {text}")
        self.badge.setStyleSheet(f"color:{clr}; font-weight:bold;")

    def _on_position_changed(self, pos):
        dur = self.player.duration() or 0
        self.time_lbl.setText(f"{fmt_time(pos)} / {fmt_time(dur)}")

    def _on_duration_changed(self, dur):
        if dur:
            self.end_spin.setValue(dur // 1000)

    # ── Processing ─────────────────────────────────────────────────────────────
    def _start_processing(self):
        if not self.video_path:
            QMessageBox.warning(self,"No Video","Please load a video file first."); return
        if not self.questions:
            QMessageBox.warning(self,"No Questions","Please load a questions file first."); return
        self.range_start = self.start_spin.value() * 1000
        self.range_end   = self.end_spin.value()   * 1000
        self.step_ms     = int(self.step_spin.value() * 1000)
        self.label       = self.label_combo.currentText()
        if self.range_end <= self.range_start:
            QMessageBox.warning(self,"Invalid Range","End time must be greater than start time."); return
        self.processing   = True
        self.total_steps  = max(1,(self.range_end - self.range_start) // self.step_ms)
        self.current_pos  = self.range_start
        self.btn_start.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.player.setPosition(self.current_pos)
        self.player.play()
        QTimer.singleShot(500, self._begin_step)

    def _begin_step(self):
        if not self.processing: return
        self.player.pause()
        self.current_q   = 0
        self.answers_buf = {}
        step_num = (self.current_pos - self.range_start) // self.step_ms + 1
        self._set_status(f"Step {step_num} / {self.total_steps}  •  {fmt_time(self.current_pos)}")
        self._show_question()

    def _show_question(self):
        if self.current_q >= len(self.questions):
            self._save_step_records()
            self._advance_step()
            return
        self.q_counter.setText(f"Question {self.current_q+1} / {len(self.questions)}")
        self.q_text.setText(self.questions[self.current_q])
        self.ans_input.clear()
        self.ans_input.setEnabled(True)
        self.btn_next_q.setEnabled(True)
        self.btn_skip_q.setEnabled(True)
        self.waiting_ans = True

    def _next_question(self):
        if not self.waiting_ans: return
        self.answers_buf[self.current_q] = self.ans_input.toPlainText().strip()
        self.current_q += 1
        self.ans_input.setEnabled(False); self.btn_next_q.setEnabled(False); self.btn_skip_q.setEnabled(False)
        self.waiting_ans = False
        self._show_question()

    def _skip_question(self):
        self.answers_buf[self.current_q] = ""
        self.current_q += 1
        self.ans_input.setEnabled(False); self.btn_next_q.setEnabled(False); self.btn_skip_q.setEnabled(False)
        self.waiting_ans = False
        self._show_question()

    def _save_step_records(self):
        ts = fmt_time(self.current_pos)
        step_num = (self.current_pos - self.range_start) // self.step_ms + 1
        vname = os.path.basename(self.video_path)
        for qi, q in enumerate(self.questions):
            rec = {"video":vname,"label":self.label,"timestamp":ts,"step":step_num,"question":q,"answer":self.answers_buf.get(qi,"")}
            self.records.append(rec)
            self._add_table_row(rec)
        self.btn_export.setEnabled(True)

    def _add_table_row(self, rec):
        row = self.table.rowCount()
        self.table.insertRow(row)
        clr = QColor(MOTHER_CLR) if rec["label"]=="Mother" else QColor(CHILD_CLR)
        for col, val in enumerate([rec["video"],rec["label"],rec["timestamp"],str(rec["step"]),rec["question"],rec["answer"]]):
            item = QTableWidgetItem(val)
            item.setForeground(clr)
            self.table.setItem(row, col, item)
        self.table.scrollToBottom()

    def _advance_step(self):
        self.current_pos += self.step_ms
        if self.current_pos >= self.range_end:
            self._finish_processing(); return
        self.player.setPosition(self.current_pos)
        self.player.play()
        QTimer.singleShot(500, self._begin_step)

    def _finish_processing(self):
        self.processing = False
        self.player.pause()
        self.btn_start.setEnabled(True); self.btn_stop.setEnabled(False)
        self.ans_input.setEnabled(False); self.btn_next_q.setEnabled(False); self.btn_skip_q.setEnabled(False)
        self.q_text.setText("(processing complete)")
        self.q_counter.setText("Done!")
        self._set_status(f"✅ Finished! {len(self.records)} records. Ready to export.")
        QMessageBox.information(self,"Done",f"Processing complete!\n{len(self.records)} annotation records collected.")

    def _stop_processing(self):
        if not self.processing: return
        self.processing = False; self.waiting_ans = False
        self.player.pause()
        self.btn_start.setEnabled(True); self.btn_stop.setEnabled(False)
        self.ans_input.setEnabled(False); self.btn_next_q.setEnabled(False); self.btn_skip_q.setEnabled(False)
        self._set_status("⏹ Processing stopped.")

    # ── Export ─────────────────────────────────────────────────────────────────
    def _export_excel(self):
        if not self.records:
            QMessageBox.warning(self,"No Data","No records to export."); return
        path, _ = QFileDialog.getSaveFileName(self,"Save Excel","annotations.xlsx","Excel Files (*.xlsx)")
        if not path: return
        if OPENPYXL_AVAILABLE:
            self._export_openpyxl(path)
        else:
            path2 = path.replace(".xlsx",".csv")
            with open(path2,"w",newline="",encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=["video","label","timestamp","step","question","answer"])
                w.writeheader(); w.writerows(self.records)
            QMessageBox.information(self,"Exported (CSV)",f"openpyxl not found — saved as CSV:\n{path2}")

    def _export_openpyxl(self, path):
        wb = Workbook(); ws = wb.active; ws.title = "Annotations"
        headers = ["Video Title","Label","Timestamp","Step #","Question","Answer"]
        hfill = PatternFill("solid", fgColor="35355A")
        hfont = Font(bold=True, color="E0E0F0", size=12)
        bs = Side(style="thin", color="555577")
        brd = Border(left=bs,right=bs,top=bs,bottom=bs)
        for c,h in enumerate(headers,1):
            cell = ws.cell(1,c,h); cell.font=hfont; cell.fill=hfill
            cell.alignment=Alignment(horizontal="center",vertical="center"); cell.border=brd
        ws.row_dimensions[1].height = 22
        for r,rec in enumerate(self.records,2):
            is_mother = rec["label"]=="Mother"
            fill = PatternFill("solid",fgColor="3D2A10" if is_mother else "0D2E2A")
            lclr = "F4A261" if is_mother else "56CFB2"
            for c,val in enumerate([rec["video"],rec["label"],rec["timestamp"],rec["step"],rec["question"],rec["answer"]],1):
                cell = ws.cell(r,c,val); cell.fill=fill; cell.border=brd
                cell.alignment=Alignment(vertical="center",wrap_text=True)
                if c==2: cell.font=Font(bold=True,color=lclr)
        for i,w in zip(range(1,7),[35,10,12,8,55,65]):
            ws.column_dimensions[ws.cell(1,i).column_letter].width=w
        ws.freeze_panes="A2"
        wb.save(path)
        self._set_status(f"✅ Exported: {os.path.basename(path)}")
        QMessageBox.information(self,"Exported",f"Saved to:\n{path}")

    def _clear_table(self):
        if self.records and QMessageBox.question(self,"Clear?","Clear all records?",
            QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No) == QMessageBox.StandardButton.Yes:
            self.records.clear(); self.table.setRowCount(0); self.btn_export.setEnabled(False)
            self._set_status("Table cleared.")

    def _set_status(self, msg):
        self.status_lbl.setText(msg)


def main():
    app = QApplication(sys.argv)
    app.setApplicationName("Video Annotator")
    win = VideoAnnotator()
    win.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()